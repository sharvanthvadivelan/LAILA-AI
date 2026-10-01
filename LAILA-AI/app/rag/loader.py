import json, zipfile
from pathlib import Path

SUPPORTED = {".pdf", ".docx", ".txt", ".md", ".csv", ".json", ".xlsx"}


def extract(path):
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError("Unsupported document type.")
    if suffix in {".docx", ".xlsx"}:
        with zipfile.ZipFile(path) as z:
            if (
                sum(x.file_size for x in z.infolist()) > 50_000_000
                or len(z.infolist()) > 10000
            ):
                raise ValueError("Expanded document is too large.")
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        if len(reader.pages) > 500:
            raise ValueError("PDF limit is 500 pages.")
        output = [
            (i + 1, f"Page {i+1}", p.extract_text() or "")
            for i, p in enumerate(reader.pages)
        ]
    elif suffix == ".docx":
        from docx import Document

        doc = Document(path)
        text = "\n".join(p.text for p in doc.paragraphs)
        text += "\n" + "\n".join(
            " | ".join(c.text for c in r.cells) for t in doc.tables for r in t.rows
        )
        output = [(None, "Document text (pagination unavailable)", text)]
    elif suffix == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        output = []
        try:
            for sheet in wb:
                lines = []
                for i, row in enumerate(sheet.iter_rows(values_only=True)):
                    if i >= 10000:
                        raise ValueError("Worksheet row limit is 10,000.")
                    lines.append(" | ".join("" if x is None else str(x) for x in row))
                output.append((None, "Sheet " + sheet.title, "\n".join(lines)))
        finally:
            wb.close()
    else:
        text = path.read_text(encoding="utf-8-sig")
        if suffix == ".json":
            text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
        output = [(None, "Text", text)]
    if sum(len(t) for _, _, t in output) > 2_000_000:
        raise ValueError("Extracted text exceeds 2 million characters.")
    if not any(t.strip() for _, _, t in output):
        raise ValueError("No extractable text. Scanned PDFs need OCR first.")
    return output
