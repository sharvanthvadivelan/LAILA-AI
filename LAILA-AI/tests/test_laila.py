import asyncio, base64, io, json
from pathlib import Path
import httpx, pytest
from app.database import database as db
from app.core.security import safe_path, validate_root
from app.tools.calculator import calculate
from app.llm.ollama import OllamaProvider
from app.llm.base import LLMError
from app.rag.loader import extract
from app.tools.registry import ToolRegistry
from app.agent.orchestrator import Agent
from app.core.markdown import render_markdown
from app.vision.vision import validate_image
from app.main import app


class FakeOllama:
    """TEST ONLY; production never imports fixtures."""

    def __init__(self, tools=False):
        self.calls = 0
        self.tools = tools
        self.messages = []

    async def show(self, model):
        return {"capabilities": ["tools"] if self.tools else []}

    async def embed(self, model, texts):
        return [[1.0, 0.0] if "banker" in t.lower() else [0.0, 1.0] for t in texts]

    async def stream(self, model, messages, options, tools=None):
        self.calls += 1
        self.messages = messages
        if self.tools and self.calls == 1:
            yield {
                "message": {
                    "tool_calls": [
                        {
                            "function": {
                                "name": "calculator",
                                "arguments": {"expression": "7*9"},
                            }
                        }
                    ]
                },
                "done": True,
            }
        else:
            yield {
                "message": {"thinking": "PRIVATE THOUGHT", "content": "Checked "},
                "done": False,
            }
            yield {"message": {"content": "answer."}, "done": True}


def install(fake):
    app.state.provider = fake
    app.state.registry = ToolRegistry(fake)
    app.state.agent = Agent(fake, app.state.registry)


def chat(c):
    return c.post("/api/conversations").json()["id"]


def test_calculator():
    assert calculate("(25/100)*840")["result"] == 210
    assert calculate("458 × 92")["result"] == 42136
    for expression in [
        '__import__("os")',
        "2**999999",
        "(1).__class__",
        "[1,2]",
        "(-1)**0.5",
    ]:
        with pytest.raises((ValueError, SyntaxError)):
            calculate(expression)


def test_file_boundaries(tmp_path):
    root = tmp_path / "approved"
    root.mkdir()
    good = root / "n.txt"
    good.write_text("ok")
    secret = root / ".env"
    secret.write_text("secret")
    outside = tmp_path / "outside.txt"
    outside.write_text("no")
    assert safe_path(str(good), [str(root)]) == good
    for p in [secret, outside, root / ".." / "outside.txt"]:
        with pytest.raises(ValueError):
            safe_path(str(p), [str(root)])
    try:
        (root / "link.txt").symlink_to(outside)
    except OSError:
        pass
    else:
        with pytest.raises(ValueError):
            safe_path(str(root / "link.txt"), [str(root)])
    with pytest.raises(ValueError):
        validate_root(Path.home())


def test_http_security(client):
    assert (
        client.post("/api/conversations", headers={"X-Laila-Token": "bad"}).status_code
        == 403
    )
    assert (
        client.get("/api/settings", headers={"Origin": "https://evil.test"}).status_code
        == 403
    )
    assert client.get("/api/session", headers={"Host": "evil.test"}).status_code == 403
    assert (
        client.get("/api/session", headers={"Sec-Fetch-Site": "cross-site"}).status_code
        == 403
    )
    assert (
        client.post(
            "/api/conversations", headers={"Content-Length": "20000001"}
        ).status_code
        == 413
    )
    assert "connect-src 'self'" in client.get("/").headers["content-security-policy"]


def test_chat_crud_stream_edit(client):
    cid = chat(client)
    r = client.post(
        f"/api/conversations/{cid}/chat", json={"text": "Calculate 25% of 840."}
    )
    assert r.status_code == 200 and "210" in r.text, r.text
    msgs = client.get("/api/conversations/" + cid).json()["messages"]
    assert len(msgs) == 2 and msgs[-1]["meta"]["status"] == "complete"
    assert (
        client.patch(
            "/api/conversations/" + cid, json={"title": "Math", "pinned": True}
        ).status_code
        == 200
    )
    assert client.get("/api/conversations?q=Math").json()[0]["pinned"] == 1
    assert (
        client.patch(
            f'/api/conversations/{cid}/messages/{msgs[0]["id"]}',
            json={"content": "Calculate 4*5"},
        ).status_code
        == 200
    )
    assert len(client.get("/api/conversations/" + cid).json()["messages"]) == 1
    assert (
        "20"
        in client.post(
            f"/api/conversations/{cid}/chat", json={"action": "regenerate"}
        ).text
    )
    assert client.delete("/api/conversations/" + cid).status_code == 200
    assert not db.rows("SELECT * FROM messages")


@pytest.mark.asyncio
async def test_ollama_protocol():
    async def handle(req):
        body = json.loads(req.content) if req.content else {}
        if req.url.path == "/api/tags":
            return httpx.Response(200, json={"models": [{"name": "local:latest"}]})
        if req.url.path == "/api/show":
            return httpx.Response(200, json={"capabilities": ["completion"]})
        if req.url.path == "/api/embed":
            return httpx.Response(200, json={"embeddings": [[1, 0]]})
        if body.get("stream"):
            return httpx.Response(
                200,
                text='{"message":{"content":"Hi"},"done":false}\n{"message":{},"done":true}\n',
            )
        return httpx.Response(200, json={"message": {"content": "Hello"}})

    p = OllamaProvider(transport=httpx.MockTransport(handle))
    assert (await p.models())[0]["name"] == "local:latest"
    assert await p.embed("embed", ["hello"]) == [[1, 0]]
    assert (await p.complete("local", [], {}))["message"]["content"] == "Hello"
    assert len([x async for x in p.stream("local", [], {})]) == 2


@pytest.mark.asyncio
async def test_provider_errors():
    async def fail(req):
        raise httpx.ConnectError("unavailable")

    with pytest.raises(LLMError, match="start Ollama"):
        await OllamaProvider(transport=httpx.MockTransport(fail)).models()

    async def remote(req):
        return httpx.Response(200, json={"remote_host": "https://ollama.com"})

    with pytest.raises(LLMError, match="Cloud"):
        await OllamaProvider(transport=httpx.MockTransport(remote)).show("x")


def test_agent_tools_hidden_thinking(client):
    fake = FakeOllama(True)
    install(fake)
    cid = chat(client)
    r = client.post(
        f"/api/conversations/{cid}/chat", json={"text": "What is seven times nine?"}
    )
    assert "PRIVATE THOUGHT" not in r.text and "Checked " in r.text, r.text
    assert fake.calls == 2
    assert any(m["role"] == "tool" and "63" in m["content"] for m in fake.messages)


def test_approval_exactly_once(client):
    cid = chat(client)
    reg = app.state.registry
    pending = asyncio.run(
        reg.execute("notes_create", {"title": "Study", "content": "Revise Prim"}, cid)
    )
    ident = pending["pending_approval"]
    assert not db.rows("SELECT * FROM notes")
    assert (
        client.post("/api/approvals/" + ident, json={"approve": True}).status_code
        == 200
    )
    assert len(db.rows("SELECT * FROM notes")) == 1
    assert (
        client.post("/api/approvals/" + ident, json={"approve": True}).status_code
        == 409
    )
    pending = asyncio.run(reg.execute("tasks_create", {"title": "Read"}, cid))
    assert (
        client.post(
            "/api/approvals/" + pending["pending_approval"], json={"approve": False}
        ).status_code
        == 200
    )
    assert not db.rows("SELECT * FROM tasks")


def test_rag_lifecycle(client):
    install(FakeOllama())
    r = client.post(
        "/api/documents",
        files={"file": ("OS.txt", b"Banker's algorithm checks safe states.")},
    )
    assert r.status_code == 200, r.text
    ident = r.json()["id"]
    assert client.post("/api/documents/" + ident + "/index").json()["chunks"] == 1
    results = client.post("/api/documents/search", json={"query": "banker"}).json()
    assert results[0]["filename"] == "OS.txt" and results[0]["page"] is None
    cid = chat(client)
    r = client.post(
        f"/api/conversations/{cid}/chat",
        json={"text": "Explain banker", "document_ids": [ident]},
    )
    assert "OS.txt" in r.text and "sources" in r.text, r.text
    assert client.delete("/api/documents/" + ident).status_code == 200
    assert not db.rows("SELECT * FROM document_chunks")


def test_index_failure_retains_file(client):
    class Bad(FakeOllama):
        async def embed(self, *args):
            raise LLMError("Embedding model missing")

    install(Bad())
    ident = client.post(
        "/api/documents", files={"file": ("test.txt", b"hello")}
    ).json()["id"]
    assert client.post("/api/documents/" + ident + "/index").status_code == 503
    assert client.get("/api/documents").json()[0]["status"] == "error"
    assert Path(db.rows("SELECT path FROM documents")[0]["path"]).exists()


def test_extract_formats(tmp_path):
    from docx import Document
    from openpyxl import Workbook

    doc = Document()
    doc.add_paragraph("Operating systems")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Deadlock"
    p = tmp_path / "test.docx"
    doc.save(p)
    assert "Deadlock" in extract(p)[0][2]
    wb = Workbook()
    wb.active.append(["OS", 90])
    p = tmp_path / "test.xlsx"
    wb.save(p)
    assert "OS" in extract(p)[0][2]
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(width=100, height=100)
    p = tmp_path / "empty.pdf"
    with p.open("wb") as f:
        w.write(f)
    with pytest.raises(ValueError, match="OCR"):
        extract(p)


def test_data_analysis(client):
    ident = client.post(
        "/api/documents",
        files={"file": ("scores.csv", b"team,score\nA,10\nA,20\nB,30\nB,\n")},
    ).json()["id"]
    r = client.post(
        "/api/analysis",
        json={
            "document_id": ident,
            "group_by": "team",
            "value_column": "score",
            "aggregate": "mean",
        },
    ).json()
    assert r["rows"] == 4 and r["columns"][1]["missing"] == 1
    assert r["chart"] == [{"label": "A", "value": 15.0}, {"label": "B", "value": 30.0}]
    assert r["statistics"]["score"]["mean"] == 20
    r = client.post(
        "/api/analysis",
        json={
            "document_id": ident,
            "filter_column": "team",
            "filter_value": "A",
            "sort_by": "score",
            "descending": True,
        },
    ).json()
    assert r["rows"] == 2 and r["preview"][0]["score"] == 20


def test_personal_data_crud(client):
    nid = client.post("/api/notes", json={"title": "First", "content": "Text"}).json()[
        "id"
    ]
    assert (
        client.put(
            "/api/notes/" + nid, json={"title": "Changed", "content": "More"}
        ).status_code
        == 200
    )
    assert client.get("/api/notes?q=Changed").json()[0]["id"] == nid
    tid = client.post(
        "/api/tasks", json={"title": "Read", "due_date": "2026-10-01"}
    ).json()["id"]
    assert (
        client.put(
            "/api/tasks/" + tid, json={"title": "Read", "completed": True}
        ).status_code
        == 200
    )
    assert client.get("/api/tasks").json()[0]["completed"] == 1
    client.post("/api/memories", json={"content": "Be concise"})
    s = client.get("/api/settings").json()
    s["memory_enabled"] = False
    assert client.put("/api/settings", json=s).status_code == 200
    from app.memory.manager import context

    assert (
        context(db.settings()) == ""
        and len(client.get("/api/export").json()["memories"]) == 1
    )
    client.delete("/api/memories")
    assert not client.get("/api/memories").json()


def test_rendering_security():
    h = render_markdown(
        '<script>alert(1)</script>\n\n![tracker](https://evil.test/img)\n\n[bad](javascript:alert(1))\n\n```python\nprint("OK")\n```'
    )
    assert (
        "<script>" not in h
        and "<img " not in h
        and 'href="javascript:' not in h
        and "code-lang" in h
    )


def test_vision_gating(client):
    from PIL import Image

    b = io.BytesIO()
    Image.new("RGB", (2, 2)).save(b, format="PNG")
    encoded = base64.b64encode(b.getvalue()).decode()
    assert validate_image(encoded) == encoded
    with pytest.raises(Exception):
        validate_image(base64.b64encode(b"bad").decode())
    install(FakeOllama())
    cid = chat(client)
    r = client.post(
        f"/api/conversations/{cid}/chat", json={"text": "Describe", "images": [encoded]}
    )
    assert (
        "does not support image" in r.text
        and encoded not in client.get("/api/conversations/" + cid).text
    )


@pytest.mark.asyncio
async def test_cancellation_saves_partial(client):
    cid = chat(client)
    db.message(cid, "user", "test")

    class Slow(FakeOllama):
        async def stream(self, *args, **kwargs):
            yield {"message": {"content": "Partial"}, "done": False}
            await asyncio.sleep(60)

    p = Slow()
    agent = Agent(p, ToolRegistry(p))
    received = asyncio.Event()

    async def consume():
        async for event in agent.run(cid, "test", "chat", [], [], asyncio.Event()):
            if event["type"] == "token":
                received.set()

    task = asyncio.create_task(consume())
    await asyncio.wait_for(received.wait(), 2)
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task
    last = db.rows("SELECT * FROM messages ORDER BY rowid DESC")[0]
    assert (
        last["content"] == "Partial" and json.loads(last["meta"])["status"] == "stopped"
    )


def test_missing_voice_and_invalid_file(client):
    assert (
        client.post("/api/documents", files={"file": ("a.exe", b"abc")}).status_code
        == 400
    )
    assert (
        client.post(
            "/api/voice/transcribe", files={"file": ("clip.wav", b"abc")}
        ).status_code
        == 400
    )


def test_regenerate_image_failure_does_not_delete_answer(client):
    from PIL import Image

    b = io.BytesIO()
    Image.new("RGB", (2, 2)).save(b, format="PNG")
    encoded = base64.b64encode(b.getvalue()).decode()
    install(FakeOllama())
    cid = chat(client)
    client.post(
        f"/api/conversations/{cid}/chat", json={"text": "Describe", "images": [encoded]}
    )
    before = client.get("/api/conversations/" + cid).json()["messages"]
    assert (
        client.post(
            f"/api/conversations/{cid}/chat", json={"action": "regenerate"}
        ).status_code
        == 400
    )
    assert client.get("/api/conversations/" + cid).json()["messages"] == before


def test_pdf_actual_page_metadata(tmp_path):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

    w = PdfWriter()
    page = w.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): w._add_object(font)}
            )
        }
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 72 720 Td (Banker algorithm) Tj ET")
    page[NameObject("/Contents")] = w._add_object(stream)
    p = tmp_path / "notes.pdf"
    with p.open("wb") as f:
        w.write(f)
    sections = extract(p)
    assert (
        sections[0][0] == 1
        and sections[0][1] == "Page 1"
        and "Banker" in sections[0][2]
    )


def test_missing_approved_folder_does_not_break_settings(client, tmp_path):
    directory = tmp_path / "workspace"
    directory.mkdir()
    s = client.get("/api/settings").json()
    s["approved_directories"] = [str(directory)]
    assert client.put("/api/settings", json=s).status_code == 200
    directory.rmdir()
    assert client.get("/api/settings").status_code == 200
    with pytest.raises((OSError, ValueError)):
        safe_path(str(directory), [str(directory)])


def test_tools_disabled_and_data_capability(client):
    fake = FakeOllama()
    install(fake)
    s = client.get("/api/settings").json()
    s["tools_enabled"] = False
    client.put("/api/settings", json=s)
    cid = chat(client)
    r = client.post(f"/api/conversations/{cid}/chat", json={"text": "Calculate 4*5"})
    assert fake.calls == 1 and not db.rows("SELECT * FROM tool_logs")
    r = client.post(
        f"/api/conversations/{cid}/chat",
        json={"text": "Analyze my data", "mode": "data"},
    )
    assert "needs a local model with tool support" in r.text
